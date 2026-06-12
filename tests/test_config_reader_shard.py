import importlib.util
import sys
import types
from pathlib import Path

from openpyxl import Workbook


def _load_config_reader_module():
    scripts_dir = Path(__file__).parents[1] / "knowledge" / "scripts"
    sys.path.insert(0, str(scripts_dir))
    spec = importlib.util.spec_from_file_location("hui_test_config_reader", scripts_dir / "config_reader.py")
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def test_compute_shard_uses_last8_in_client_no_for_business_db_and_table(monkeypatch, tmp_path):
    module = _load_config_reader_module()

    class FakeCursor:
        def execute(self, _sql, _params):
            return None

        def fetchone(self):
            return ("33060210021113",)

    class FakeConnection:
        def cursor(self):
            return FakeCursor()

        def close(self):
            return None

    fake_pymysql = types.SimpleNamespace(connect=lambda **_kwargs: FakeConnection())
    monkeypatch.setitem(sys.modules, "pymysql", fake_pymysql)

    fake_config = types.SimpleNamespace(
        env_dbs={
            "ifm60-dxfund": {
                "databases": {
                    "ui_bc": {"user": "u", "password": "p"},
                    "ui_fund1": {},
                    "ui_fund2": {},
                }
            }
        },
        get_db_config=lambda _env: {
            "common": {"host": "127.0.0.1", "port": 3306},
            "databases": {
                "ui_bc": {"user": "u", "password": "p", "database": "ui_bc"},
                "ui_fund1": {},
                "ui_fund2": {},
            },
        },
    )
    conf_module = types.ModuleType("conf")
    config_module = types.ModuleType("conf.config")
    config_module.Config = fake_config
    monkeypatch.setitem(sys.modules, "conf", conf_module)
    monkeypatch.setitem(sys.modules, "conf.config", config_module)

    reader = module.ConfigReader.__new__(module.ConfigReader)
    reader.get_env = lambda key, default=None: str(tmp_path) if key == "project_dir" else (default or "")
    reader.get_preset = lambda key, default="": "A33060210021113" if key == "customer_id" else default

    result = reader.compute_shard("dxfund")

    assert result["db_name"] == "ui_fund2"
    assert result["trans_db_num"] == 2
    assert result["trans_table_num"] == 13
    assert result["pub_table_num"] == 10
    assert result["route_number"] == 10021113


def test_compute_shard_falls_back_to_config_in_client_no_when_db_mapping_is_blank(monkeypatch, tmp_path):
    module = _load_config_reader_module()

    class FakeCursor:
        def execute(self, _sql, _params):
            return None

        def fetchone(self):
            return (" ",)

    class FakeConnection:
        def cursor(self):
            return FakeCursor()

        def close(self):
            return None

    fake_pymysql = types.SimpleNamespace(connect=lambda **_kwargs: FakeConnection())
    monkeypatch.setitem(sys.modules, "pymysql", fake_pymysql)

    fake_config = types.SimpleNamespace(
        env_dbs={
            "ifm60-trust": {
                "databases": {
                    "ui_bc": {"user": "u", "password": "p"},
                    "ui_trust1": {},
                    "ui_trust2": {},
                }
            }
        },
        get_db_config=lambda _env: {
            "common": {"host": "127.0.0.1", "port": 3306},
            "databases": {
                "ui_bc": {"user": "u", "password": "p", "database": "ui_bc"},
                "ui_trust1": {},
                "ui_trust2": {},
            },
        },
    )
    conf_module = types.ModuleType("conf")
    config_module = types.ModuleType("conf.config")
    config_module.Config = fake_config
    monkeypatch.setitem(sys.modules, "conf", conf_module)
    monkeypatch.setitem(sys.modules, "conf.config", config_module)

    presets = {
        "customer_id": "A33060210021113",
        "in_client_no": "33060210021113",
    }
    reader = module.ConfigReader.__new__(module.ConfigReader)
    reader.get_env = lambda key, default=None: str(tmp_path) if key == "project_dir" else (default or "")
    reader.get_preset = lambda key, default="": presets.get(key, default)

    result = reader.compute_shard("dxtrust")

    assert result["in_client_no"] == "33060210021113"
    assert result["last8"] == "10021113"
    assert result["db_name"] == "ui_trust2"


def test_compute_shard_uses_single_db_key_without_shard_calculation(tmp_path):
    module = _load_config_reader_module()
    presets = {
        "customer_id": "A33060210021113",
        "in_client_no": "33060210021113",
    }
    env = {
        "project_dir": str(tmp_path),
        "single_db_key": "GYYHAUTO",
        "env_key": "gyyh-auto",
        "db_type": "oracle",
    }
    reader = module.ConfigReader.__new__(module.ConfigReader)
    reader.get_env = lambda key, default=None: env.get(key, default or "")
    reader.get_preset = lambda key, default="": presets.get(key, default)

    result = reader.compute_shard("dxfund")

    assert result["env_name"] == "gyyh-auto"
    assert result["db_name"] == "GYYHAUTO"
    assert result["routing_db"] == "GYYHAUTO"
    assert result["db_type"] == "oracle"
    assert result["formula"] == "single_db(GYYHAUTO)"


def test_presets_can_be_derived_from_chinese_fields_without_field_key_column(tmp_path):
    module = _load_config_reader_module()
    workbook = Workbook()
    env = workbook.active
    env.title = "环境配置"
    env.append(["配置项", "配置值", "说明"])
    env.append(["module", "dxtrust", ""])
    menus = workbook.create_sheet("菜单导入")
    menus.append(["状态", "菜单名称"])
    presets = workbook.create_sheet("预埋数据")
    presets.append(["字段名", "默认值", "说明", "适用模块"])
    presets.append(["客户标识", "A33060210021113", "统一客户标识", "公共"])
    presets.append(["TA代码", "HT", "TA代码", "信托代销"])
    presets.append(["产品代码", "UI_HT_P001", "产品信息查询用产品代码", "信托代销"])
    presets.append(["产品代码", "cycleset01", "产品信息查询用产品代码", "基金代销"])
    config_path = tmp_path / "test_config.xlsx"
    workbook.save(config_path)

    reader = module.ConfigReader(config_path)

    assert reader.get_presets("dxtrust") == {
        "customer_id": "A33060210021113",
        "ta_code": "HT",
        "prd_code": "UI_HT_P001",
    }
    assert reader.get_preset_fields("dxtrust")["customer_id"] == {
        "name": "客户标识",
        "value": "A33060210021113",
    }


def test_init_menu_prefers_excel_env_key_for_single_db_environment(tmp_path):
    module = _load_config_reader_module()
    project = tmp_path / "project"
    (project / "conf").mkdir(parents=True)
    (project / "conf" / "__init__.py").write_text("", encoding="utf-8")
    (project / "conf" / "config.py").write_text(
        "class Config:\n"
        "    TEST_CASE_DIRS = {'信托代销': 'tests_out'}\n"
        "    env_dbs = {}\n",
        encoding="utf-8",
    )

    workbook = Workbook()
    env = workbook.active
    env.title = "环境配置"
    env.append(["配置项", "配置值", "说明"])
    env.append(["project_dir", str(project), ""])
    env.append(["env_key", "gyyh-auto", ""])
    env.append(["single_db_key", "GYYHAUTO", ""])
    env.append(["db_type", "oracle", ""])
    env.append(["log_path_module_keywords", "信托:dxtrust", ""])
    menus = workbook.create_sheet("菜单导入")
    menus.append(["状态", "菜单名称", "导航一级菜单", "导航二级菜单", "导航三级菜单", "菜单URL"])
    menus.append(["未完成", "菜单A", "信托代销", "查询管理", "菜单A", "/menu/FUNC_A"])
    presets = workbook.create_sheet("预埋数据")
    presets.append(["字段名", "默认值", "说明", "适用模块"])
    presets.append(["客户标识", "A33060210021113", "", "公共"])
    config_path = tmp_path / "test_config.xlsx"
    workbook.save(config_path)

    reader = module.ConfigReader(config_path)

    result = reader.init_menu("菜单A")

    assert result["env_key"] == "gyyh-auto"
    assert result["env"]["single_db_key"] == "GYYHAUTO"
    assert result["env"]["db_type"] == "oracle"
