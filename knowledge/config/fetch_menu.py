# -*- coding: utf-8 -*-
"""Fetch menu URLs from the configured environment database.

This module is intentionally side-effect free on import. Use ``main`` or call
the functions directly from tests/scripts.
"""

import argparse
import importlib.util
import json
import os
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional


CONFIG_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG_XLSX = CONFIG_DIR / "test_config.xlsx"
DEFAULT_OUTPUT = CONFIG_DIR / "menu_urls.json"
_ORACLE_CLIENT_INITIALIZED = False


def _read_excel_env_defaults(xlsx_path: Path = DEFAULT_CONFIG_XLSX) -> Dict[str, str]:
    """Read defaults from the 环境配置 sheet when CLI args are omitted."""
    if not xlsx_path.exists():
        return {}

    try:
        from openpyxl import load_workbook
    except ImportError:
        return {}

    wb = load_workbook(xlsx_path, data_only=True)
    try:
        ws = wb["环境配置"]
        values: Dict[str, str] = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            key = str(row[0] or "").strip() if row else ""
            value = str(row[1] or "").strip() if row and row[1] is not None else ""
            if key and value:
                values[key] = value
        return values
    finally:
        wb.close()


def _load_project_config(project_dir: str):
    config_path = Path(project_dir) / "conf" / "config.py"
    if not config_path.exists():
        raise FileNotFoundError(f"找不到项目配置文件: {config_path}")

    spec = importlib.util.spec_from_file_location("hui_project_config", config_path)
    if spec is None or spec.loader is None:
        raise ImportError(f"无法加载项目配置文件: {config_path}")

    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.Config


def load_db_config(project_dir: str, env_key: str, db_key: str) -> Dict[str, Any]:
    """Load and merge one database config from project ``conf/config.py``."""
    if not project_dir:
        raise ValueError("project_dir 不能为空")
    if not env_key:
        raise ValueError("env_key 不能为空")
    if not db_key:
        raise ValueError("db_key 不能为空")

    config_cls = _load_project_config(project_dir)
    env_db = config_cls.get_db_config(env_key)
    if not env_db:
        raise KeyError(f"环境不存在: {env_key}")

    common = env_db.get("common", {})
    databases = env_db.get("databases", {})
    db_cfg = databases.get(db_key)
    if not db_cfg:
        raise KeyError(f"数据库配置不存在: {env_key}.{db_key}")

    merged = dict(common)
    merged.update(db_cfg)
    merged.setdefault("database", db_key)
    merged.setdefault("db_key", db_key)
    merged.setdefault("db_type", common.get("db_type", "mysql"))
    return merged


def connect_db(db_cfg: Dict[str, Any]):
    """Create a DB connection for Oracle or MySQL-style configs."""
    db_type = str(db_cfg.get("db_type") or "mysql").lower()
    if db_type == "oracle":
        try:
            import oracledb
        except ImportError as exc:
            raise ImportError("请安装 oracledb: pip install oracledb") from exc

        _init_oracle_client_if_available(oracledb, db_cfg)
        database = db_cfg.get("service_name") or db_cfg.get("database") or db_cfg.get("db_key")
        dsn = db_cfg.get("dsn") or f"{db_cfg['host']}:{db_cfg.get('port', 1521)}/{database}"
        return oracledb.connect(
            user=db_cfg["user"],
            password=db_cfg["password"],
            dsn=dsn,
        )

    try:
        import pymysql
    except ImportError as exc:
        raise ImportError("请安装 pymysql: pip install pymysql") from exc

    return pymysql.connect(
        host=db_cfg["host"],
        port=int(db_cfg.get("port", 3306)),
        user=db_cfg["user"],
        password=db_cfg["password"],
        database=db_cfg.get("database") or db_cfg.get("db_key"),
        charset=db_cfg.get("charset", "utf8mb4"),
        cursorclass=pymysql.cursors.DictCursor,
    )


def _init_oracle_client_if_available(oracledb, db_cfg: Dict[str, Any]) -> None:
    global _ORACLE_CLIENT_INITIALIZED
    if _ORACLE_CLIENT_INITIALIZED or not hasattr(oracledb, "init_oracle_client"):
        return

    lib_dir = _oracle_client_lib_dir(db_cfg)
    if not lib_dir:
        return

    try:
        oracledb.init_oracle_client(lib_dir=lib_dir)
        _ORACLE_CLIENT_INITIALIZED = True
    except Exception as exc:
        message = str(exc).lower()
        if "already initialized" in message or "dpy-2017" in message:
            _ORACLE_CLIENT_INITIALIZED = True
            return
        raise


def _oracle_client_lib_dir(db_cfg: Dict[str, Any]) -> str:
    for key in ("oracle_client_lib_dir", "instant_client_dir", "client_lib_dir", "lib_dir"):
        value = str(db_cfg.get(key) or "").strip()
        if value:
            return value

    for env_key in ("ORACLE_CLIENT_LIB_DIR", "ORACLE_INSTANT_CLIENT_DIR"):
        value = str(os.environ.get(env_key) or "").strip()
        if value:
            return value

    for path_item in os.environ.get("PATH", "").split(os.pathsep):
        if path_item and (Path(path_item) / "oci.dll").exists():
            return path_item
    return ""


def _row_to_dict(row: Any, columns: Iterable[str]) -> Dict[str, Any]:
    if isinstance(row, dict):
        return {str(key).lower(): value for key, value in row.items()}
    return {column.lower(): value for column, value in zip(columns, row)}


def fetch_menu_rows(conn) -> List[Dict[str, Any]]:
    """Fetch ``tsys_menu`` rows and filter out ifmcounter menus."""
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT * FROM tsys_menu")
        rows = cursor.fetchall()
        columns = [str(col[0]) for col in (cursor.description or [])]
    finally:
        cursor.close()

    filtered_rows: List[Dict[str, Any]] = []
    for row in rows:
        normalized = _row_to_dict(row, columns)
        row_text = json.dumps(normalized, ensure_ascii=False, default=str).lower()
        if "ifmcounter" not in row_text:
            filtered_rows.append(normalized)
    return filtered_rows


def write_menu_urls(rows: List[Dict[str, Any]], output_path: Path) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )


def _resolve_args(args: argparse.Namespace) -> argparse.Namespace:
    defaults = _read_excel_env_defaults()
    args.project_dir = args.project_dir or defaults.get("project_dir")
    args.env_key = args.env_key or defaults.get("env_key")
    args.db_key = args.db_key or defaults.get("single_db_key") or defaults.get("db_mapping_dbs")
    args.output = Path(args.output) if args.output else DEFAULT_OUTPUT
    return args


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="从环境数据库导出菜单 URL 配置")
    parser.add_argument("--project-dir", help="项目根目录（conf/config.py 所在目录）")
    parser.add_argument("--env", dest="env_key", help="环境标识，如 gyyh-auto")
    parser.add_argument("--db", dest="db_key", help="数据库 key，如 GYYHAUTO")
    parser.add_argument("--output", help="输出 JSON 路径，默认 knowledge/config/menu_urls.json")
    args = _resolve_args(parser.parse_args(argv))

    db_cfg = load_db_config(args.project_dir, args.env_key, args.db_key)
    conn = connect_db(db_cfg)
    try:
        rows = fetch_menu_rows(conn)
    finally:
        conn.close()

    write_menu_urls(rows, args.output)
    print(f"菜单 URL 已写入: {args.output}，共 {len(rows)} 条")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
