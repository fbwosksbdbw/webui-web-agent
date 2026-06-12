import os
import sys
import json

if hasattr(sys.stdout, 'reconfigure'):
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
from openpyxl import load_workbook
from module_registry import (
    env_keyword,
    infer_module_from_keyword_mapping,
    infer_module_from_text,
    module_to_chinese,
    normalize_module,
    sql_data_attr,
)
from skill_config import resolve_env

# 配置文件默认路径（相对于skill根目录）
SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_CONFIG_PATH = os.path.join(SKILL_DIR, 'config', 'test_config.xlsx')
FIELD_SEMANTICS_PATH = os.path.join(SKILL_DIR, 'config', 'field_semantics.json')

# 模块英文标识 → Excel「适用模块」列中文名
_MODULE_EN_TO_CN = {
    'dxfund': '基金代销',
    'dxtrust': '信托代销',
    'fina': '理财代销',
    'dxcp': '商业养老金',
    'per': '个人养老金',
    'pension': '个人养老金',
}
_MENU_NAME_PREFIXES = (
    '基金代销',
    '信托代销',
    '理财代销',
    '商业养老金',
    '个人养老金',
    '保险代销',
    '汇市通',
)


def _normalize_menu_display_name(menu_name: str) -> str:
    text = str(menu_name or '').strip()
    changed = True
    while changed:
        changed = False
        for prefix in _MENU_NAME_PREFIXES:
            if text.startswith(prefix) and len(text) > len(prefix):
                text = text[len(prefix):].strip()
                changed = True
    return text


def _menu_tree_depth(tree_idx: str) -> int:
    return len([part for part in str(tree_idx or '').split('#') if part])


def _menu_url_rank(entry: dict) -> tuple:
    tree_idx = entry.get('tree_idx', '')
    depth = _menu_tree_depth(tree_idx)
    is_tree_idx = 1 if '#' in str(tree_idx) else 0
    # 同名父子节点并存时，优先叶子/更深节点，避免勘测停在目录节点。
    return (is_tree_idx, depth)


def _load_field_semantics() -> dict:
    try:
        with open(FIELD_SEMANTICS_PATH, 'r', encoding='utf-8') as f:
            payload = json.load(f)
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def _semantic_key_for_label(label: str) -> str:
    mapping = _load_field_semantics().get('label_to_semantic', {})
    if not isinstance(mapping, dict):
        return ''
    label_text = str(label or '').strip()
    if not label_text:
        return ''
    if label_text in mapping:
        return str(mapping[label_text] or '').strip()
    for candidate, semantic in sorted(mapping.items(), key=lambda item: len(str(item[0])), reverse=True):
        candidate_text = str(candidate or '').strip()
        if candidate_text and (candidate_text in label_text or label_text in candidate_text):
            return str(semantic or '').strip()
    return ''


def _header_index(headers: list, name: str) -> int:
    return headers.index(name) if name in headers else -1


def _row_value(row: tuple, headers: list, name: str):
    index = _header_index(headers, name)
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _normalize_in_client_no(*candidates) -> str:
    """Return the first usable internal client number from DB/config/account values."""
    for value in candidates:
        text = str(value or '').strip()
        if not text:
            continue
        if text.isdigit():
            return text
        digits = ''.join(ch for ch in text if ch.isdigit())
        if digits:
            return digits
    return ''


class ConfigReader:
    """测试配置读取器。
    配置文件: config/test_config.xlsx（3个sheet：环境配置、菜单导入、预埋数据）
    """

    def __init__(self, config_path=None):
        self.config_path = config_path or DEFAULT_CONFIG_PATH
        if not os.path.exists(self.config_path):
            raise FileNotFoundError(f"配置文件不存在: {self.config_path}")
        self._wb = load_workbook(self.config_path, data_only=True)
        self._env = None
        self._import_menus = None
        self._menus = None
        self._presets = None

    # ========== Sheet1: 环境配置 ==========
    @property
    def env(self) -> dict:
        """环境配置，返回 {配置项: 配置值} 字典（仅含 Excel 中配置的项）"""
        if self._env is None:
            ws = self._wb['环境配置']
            self._env = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] and row[1]:  # 配置项 + 配置值
                    self._env[str(row[0])] = str(row[1])
        return self._env

    def get_env(self, key: str, default=None) -> str:
        return self.env.get(key, default or '')

    @staticmethod
    def _compute_log_path(module: str, env_data: dict) -> str:
        """根据 module 拼接完整日志路径（log_path_{module} + log_file）。"""
        log_path_key = f"log_path_{module}"
        log_dir = env_data.get(log_path_key, '')
        log_file = env_data.get('log_file', '')
        if log_dir and log_file:
            return log_dir + log_file
        return ''

    # ========== Sheet2: 菜单导入 ==========
    @property
    def import_menus(self) -> list:
        """菜单列表（仅从「菜单导入」sheet 读取，去重，过滤已完成）"""
        if self._import_menus is None:
            ws = self._wb['菜单导入']
            headers = [cell.value for cell in ws[1]]
            seen = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    d = {k: (str(v) if v is not None else '') for k, v in zip(headers, row) if k is not None}
                    name = d.get('菜单名称', '')
                    # 状态过滤：空值视为「未完成」，已完成的跳过
                    status = d.get('状态', '')
                    if status == '已完成':
                        continue
                    if name and name not in seen:
                        seen[name] = d
            self._import_menus = list(seen.values())
        return self._import_menus

    @property
    def menus(self) -> list:
        """菜单配置列表（同 import_menus，保留接口兼容）"""
        if self._menus is None:
            self._menus = self.import_menus
        return self._menus

    def get_menu(self, menu_name: str) -> dict:
        """按菜单名称获取配置（仅未完成菜单）"""
        for m in self.menus:
            if m.get('菜单名称') == menu_name:
                return m
        return {}

    def _get_menu_all(self, menu_name: str) -> dict:
        """按菜单名称获取配置（含已完成菜单，用于状态校验）"""
        ws = self._wb['菜单导入']
        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                d = {k: (str(v) if v is not None else '') for k, v in zip(headers, row) if k is not None}
                if d.get('菜单名称') == menu_name:
                    return d
        return {}

    @staticmethod
    def _extract_func_code(url: str) -> str:
        """从菜单URL末段提取功能代码。
        如 /console-dxtrust-vue/a/b/c/ → c
        """
        if not url:
            return ''
        if '#' in url:
            parts = [p for p in url.strip('#').split('#') if p]
            return parts[-1] if parts else ''
        parts = [p for p in url.rstrip('/').split('/') if p]
        return parts[-1] if parts else ''

    def get_func_code(self, menu_name: str) -> str:
        """从菜单URL动态获取功能代码"""
        url = self.get_menu_url(menu_name)
        return self._extract_func_code(url)

    def get_menu_by_code(self, func_code: str) -> dict:
        """按功能代码获取配置（从URL动态匹配）"""
        for m in self.menus:
            name = m.get('菜单名称', '')
            if self.get_func_code(name) == func_code:
                return m
        return {}

    # ========== Sheet3: 预埋数据 ==========
    @property
    def presets(self) -> dict:
        """预埋数据（全部，不过滤）。向后兼容。"""
        if self._presets is None:
            self._presets = self.get_presets()
        return self._presets

    @property
    def preset_fields(self) -> dict:
        """预埋数据完整信息，按当前模块过滤。"""
        if not hasattr(self, '_preset_fields'):
            # 🔧 修复：传入 module 参数，按模块过滤 preset_fields
            module = self.get_preset('module', '') or self.env.get('module', '')
            self._preset_fields = self.get_preset_fields(module)
        return self._preset_fields

    def get_preset(self, key: str, default=None) -> str:
        return self.presets.get(key, default or '')

    def get_presets(self, module: str = '') -> dict:
        """预埋数据，按模块过滤。module 为英文标识(dxfund/dxtrust/fina/dxcp)，为空时返回全部。"""
        ws = self._wb['预埋数据']
        headers = [cell.value for cell in ws[1]]
        module_idx = _header_index(headers, '适用模块')
        module_cn = module_to_chinese(module)
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = _row_value(row, headers, '字段名')
            value = _row_value(row, headers, '默认值')
            if not label or value is None or str(value).strip() == '':
                continue
            if module and module_idx >= 0:
                row_module = str(row[module_idx] or '').strip()
                if row_module and row_module not in ('公共', '所有菜单', module_cn):
                    continue
            key = _row_value(row, headers, '字段键') or _semantic_key_for_label(str(label)) or str(label)
            result[str(key)] = str(value)
        return result

    def get_preset_fields(self, module: str = '') -> dict:
        """预埋数据完整信息，按模块过滤。module 为英文标识。"""
        ws = self._wb['预埋数据']
        headers = [cell.value for cell in ws[1]]
        module_idx = _header_index(headers, '适用模块')
        module_cn = module_to_chinese(module)
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = _row_value(row, headers, '字段名')
            value = _row_value(row, headers, '默认值')
            if not label or value is None or str(value).strip() == '':
                continue
            if module and module_idx >= 0:
                row_module = str(row[module_idx] or '').strip()
                if row_module and row_module not in ('公共', '所有菜单', module_cn):
                    continue
            key = _row_value(row, headers, '字段键') or _semantic_key_for_label(str(label)) or str(label)
            result[str(key)] = {
                'name': str(label),
                'value': str(value),
            }
        return result

    # ========== 菜单 URL 缓存（menu_urls.json，独立于 Excel） ==========
    @property
    def menu_urls(self) -> dict:
        """{menu_code: {menu_name, tree_idx}}"""
        if not hasattr(self, '_menu_urls'):
            url_file = os.path.join(os.path.dirname(self.config_path), 'menu_urls.json')
            if os.path.exists(url_file):
                with open(url_file, 'r', encoding='utf-8') as f:
                    data = json.load(f)
                # menu_urls.json may be a list of objects or a dict keyed by menu_code
                if isinstance(data, list):
                    self._menu_urls = {item['menu_code']: item for item in data if 'menu_code' in item}
                else:
                    self._menu_urls = data
            else:
                self._menu_urls = {}
        return self._menu_urls

    def get_menu_url_by_code(self, menu_code: str) -> str:
        entry = self.menu_urls.get(menu_code, {})
        return entry.get('tree_idx', '')

    def get_menu_url_by_name(self, menu_name: str) -> str:
        normalized_name = _normalize_menu_display_name(menu_name)
        exact_matches = [
            entry for entry in self.menu_urls.values()
            if entry.get('menu_name') == menu_name and entry.get('tree_idx')
        ]
        if exact_matches:
            return max(exact_matches, key=_menu_url_rank).get('tree_idx', '')

        normalized_matches = []
        for entry in self.menu_urls.values():
            entry_name = entry.get('menu_name', '')
            if entry_name and _normalize_menu_display_name(entry_name) == normalized_name and entry.get('tree_idx'):
                normalized_matches.append(entry)
        if normalized_matches:
            return max(normalized_matches, key=_menu_url_rank).get('tree_idx', '')
        return ''

    def get_menu_url(self, menu_name: str) -> str:
        """获取菜单 URL（优先从 menu_urls.json，回退到 Excel）"""
        url = self.get_menu_url_by_name(menu_name)
        if url:
            return url
        menu = self.get_menu(menu_name)
        return menu.get('菜单URL', '')

    # ========== 基础配置表缓存（base_tables.json） ==========
    @property
    def base_tables(self) -> dict:
        if not hasattr(self, '_base_tables'):
            bt_file = os.path.join(os.path.dirname(self.config_path), 'base_tables.json')
            if os.path.exists(bt_file):
                with open(bt_file, 'r', encoding='utf-8') as f:
                    self._base_tables = json.load(f)
            else:
                self._base_tables = {}
        return self._base_tables

    def get_ta_record(self, ta_code: str, prd_type: str) -> dict:
        key = f"{ta_code}_prd{prd_type}"
        records = self.base_tables.get('tbtainfo', {}).get('records', {})
        return records.get(key, {})

    def get_product_record(self, module: str, key: str) -> dict:
        records = self.base_tables.get(module, {}).get('records', {})
        return records.get(key, {})

    def get_products_by_module(self, module: str) -> dict:
        return self.base_tables.get(module, {}).get('records', {})

    def get_client_mapping(self, bank_acc: str) -> dict:
        records = self.base_tables.get('tbclientseller2', {}).get('records', {})
        return records.get(bank_acc, {})

    # ========== 探勘字段匹配 ==========

    def match_presets(self, explored_fields: list) -> dict:
        """将探勘字段与 Excel 预埋数据的「字段名」列模糊匹配，返回 form_fields.all 结构。

        委托给 explore_wrapper.generate_field_values() 避免重复实现。
        """
        from stagehand_explorer.src.explore_wrapper import generate_field_values
        return generate_field_values(explored_fields, self.preset_fields)

    # ========== 代码生成辅助方法 ==========
    def gen_navigate_code(self, menu_name: str) -> str:
        """生成 _navigate 函数代码"""
        url = self.get_menu_url(menu_name)
        if not url:
            return f"# 未找到菜单URL: {menu_name}"
        return f'''async def _navigate(page):
    base = page.url.split('#')[0]
    await page.goto(f"{{base}}#{url}", wait_until="domcontentloaded")
    await page.wait_for_timeout(3000)
    await _collapse_sidebar(page)'''

    def gen_hierarchical_case(self, menu_name: str, test_type: str) -> str:
        """生成 hierarchical_case 装饰器"""
        menu = self.get_menu(menu_name)
        l1 = menu.get('导航一级菜单', '')
        l2 = menu.get('导航二级菜单', '')
        code = self.get_func_code(menu_name)
        return f'@hierarchical_case("{l1}", "{l2}", "{menu_name}", "{menu_name}_{test_type}")'

    def gen_preset_sql_code(self, menu_name: str) -> str:
        """生成预埋 SQL 路径代码（业务表 + 基础表）"""
        menu = self.get_menu(menu_name)
        sql_file = menu.get('预埋SQL文件名', '')
        func_code = self.get_func_code(menu_name)
        l2 = menu.get('导航二级菜单', '查询管理')
        l1 = menu.get('导航一级菜单', '')
        module_var = infer_module_from_text(l1)
        sql_attr = sql_data_attr(module_var)
        base_file = f'{func_code}_base.sql' if func_code else 'base_preset.sql'
        return f'''PRESET_SQL = os.path.join(
    Config.{sql_attr}, '{l2}', '{menu_name}', '{sql_file}'
)
PRESET_SQL_BASE = os.path.join(
    Config.{sql_attr}, '{l2}', '{menu_name}', '{base_file}'
)'''

    # ========== 分片计算 ==========
    def compute_shard(self, module: str) -> dict:
        """根据 customer_id 计算分片库（分表以日志物理表名为准，不再计算）。
        module: dxtrust | dxfund | fina | dxcp
        """
        import re

        module = normalize_module(module)
        single_db_key = str(self.get_env('single_db_key', '') or '').strip()
        if single_db_key:
            in_client_no_text = _normalize_in_client_no(
                '',
                self.get_preset('in_client_no', ''),
                self.get_preset('customer_id', ''),
            )
            return {
                'env_name': self.get_env('env_key', ''),
                'module': module,
                'customer_id': self.get_preset('customer_id', ''),
                'in_client_no': in_client_no_text,
                'last8': in_client_no_text[-8:] if len(in_client_no_text) >= 8 else in_client_no_text.zfill(8),
                'route_number': int(in_client_no_text[-8:]) if len(in_client_no_text) >= 8 and in_client_no_text[-8:].isdigit() else 0,
                'has_db2': False,
                'db_prefix': single_db_key,
                'trans_db_num': 1,
                'trans_db_remainder': 0,
                'trans_db_quotient': 0,
                'trans_table_num': 0,
                'pub_table_num': 0,
                'db_name': single_db_key,
                'routing_db': single_db_key,
                'db_type': str(self.get_env('db_type', 'mysql') or 'mysql'),
                'formula': f'single_db({single_db_key})',
            }

        import pymysql

        # 从 conf/config.py 获取数据库配置
        project_root = self.get_env('project_dir', '') or os.path.dirname(os.path.dirname(os.path.dirname(SKILL_DIR)))
        sys.path.insert(0, project_root)
        from conf.config import Config

        # module → 业务关键词（最小映射，无法从缩写推导）
        try:
            keyword = env_keyword(module)
        except ValueError as exc:
            raise ValueError(f"未知 module: {module}，支持: dxtrust/dxfund/fina/dxcp/hstc/insure/per") from exc

        # 从 Config.env_dbs 搜索匹配业务关键词的环境
        env_key = None
        biz_prefix = None
        for candidate_env, cfg in Config.env_dbs.items():
            databases = cfg.get('databases', {})
            for db_name in databases:
                if db_name != 'ui_bc' and keyword in db_name.lower():
                    env_key = candidate_env
                    m = re.match(r'^(.+?)\d+$', db_name)
                    if m:
                        biz_prefix = m.group(1)
                    break
            if env_key:
                break

        if not env_key:
            raise ValueError(
                f"Config.env_dbs 中未找到匹配 module={module} (keyword={keyword}) 的环境。"
                f"可用环境: {list(Config.env_dbs.keys())}"
            )

        db_config = Config.get_db_config(env_key)
        common = db_config.get('common', {})
        host = common.get('host')
        port = common.get('port', 3306)
        if not host:
            raise ValueError(f"环境 {env_key} 的数据库配置缺少 host")

        databases = db_config.get('databases', {})
        bc = databases.get('ui_bc', {})
        if not bc:
            raise ValueError(f"环境 {env_key} 的数据库配置缺少 ui_bc")

        customer_id = self.get_preset('customer_id', '')
        if not customer_id:
            raise ValueError("预埋数据中缺少 customer_id，请在 Excel「预埋数据」sheet 中配置")

        conn = pymysql.connect(
            host=host, port=port,
            user=bc['user'],
            password=bc['password'],
            database=bc.get('database', 'ui_bc'),
            autocommit=True
        )
        try:
            cur = conn.cursor()
            cur.execute(
                "SELECT in_client_no FROM tbclientmap WHERE bank_acc = %s AND bank_no = '000'",
                (customer_id,)
            )
            row = cur.fetchone()
            in_client_no = row[0] if row else ''
        finally:
            conn.close()

        in_client_no_text = _normalize_in_client_no(
            in_client_no,
            self.get_preset('in_client_no', ''),
            customer_id,
        )
        if not in_client_no_text:
            raise ValueError(
                f"无法计算分片：tbclientmap 中未找到有效 bank_acc={customer_id} 映射，"
                "且预埋数据中缺少有效 in_client_no/customer_id"
            )
        if len(in_client_no_text) >= 8:
            last8 = in_client_no_text[-8:]
        else:
            last8 = in_client_no_text.zfill(8)
        n = int(last8)
        pub_table_num = n % 16 + 1

        db2_name = f'{biz_prefix}2'
        has_db2 = db2_name in databases

        if has_db2:
            quotient = n // 2
            remainder = n % 2
            db_num = remainder + 1
            db_name = f'{biz_prefix}{db_num}'
            trans_table_num = quotient % 16 + 1
            formula_detail = (
                f'{n} / 2 = {quotient} 余 {remainder} -> '
                f'库{db_num}({db_name}); '
                f'{quotient} % 16 + 1 = {trans_table_num} 表'
            )
        else:
            quotient = n
            remainder = 0
            db_num = 1
            db_name = f'{biz_prefix}1'
            trans_table_num = n % 16 + 1
            formula_detail = f'{n} % 16 + 1 = {trans_table_num} 表 (单库 {db_name})'

        return {
            'env_name': env_key,
            'module': module,
            'customer_id': customer_id,
            'in_client_no': in_client_no_text,
            'last8': last8,
            'route_number': n,
            'has_db2': has_db2,
            'db_prefix': biz_prefix,
            'trans_db_num': db_num,
            'trans_db_remainder': remainder,
            'trans_db_quotient': quotient,
            'trans_table_num': trans_table_num,
            'pub_table_num': pub_table_num,
            'db_name': db_name,
            'formula': formula_detail,
        }

    def init_menu(self, menu_name: str) -> dict:
        """一次性返回一个菜单所需的全部配置（env + presets + menu + url + paths）。"""
        # 从全部菜单（含已完成）中查找，单独校验状态
        menu = self._get_menu_all(menu_name)
        if not menu:
            raise ValueError(f"「菜单导入」sheet 中未找到菜单: {menu_name}")
        status = menu.get('状态', '')
        if status == '已完成':
            raise ValueError(f"该菜单已完成，跳过: {menu_name}")
        url = self.get_menu_url(menu_name)
        func_code = self._extract_func_code(url)

        # 路径计算：project_dir（Excel环境配置） + 一级菜单 → 查 Config 映射表
        project_dir = self.get_env('project_dir', '')
        if project_dir.startswith('<') or '请配置' in project_dir:
            raise ValueError(
                f"project_dir 未配置（当前值: {project_dir}），"
                f"请在 Excel「环境配置」sheet 的 project_dir 行填入您的项目根目录"
            )
        level1 = menu.get('导航一级菜单', '')

        # 根据导航文本确定业务模块（从 Excel「环境配置」的 log_path_module_keywords 读取）。
        # GYYH 单库环境的 Excel 导航一级菜单是真实页面树（如「查询管理」），
        # 不能再把它当作旧业务模块名；查询管理默认落到理财代销工程路径。
        # 格式: "关键词1:模块名1,关键词2:模块名2"，如 "基金:dxfund,信托:dxtrust,商业养老金:dxcp"
        log_path_module_keywords = self.get_env('log_path_module_keywords', '')
        level2 = menu.get('导航二级菜单', '')
        level3 = menu.get('导航三级菜单', '')
        default_module = 'fina' if self.get_env('single_db_key', '') and str(url).startswith('ifmquerymanage#') else 'dxtrust'
        module_text = ' '.join([level1, level2, level3, menu_name, url])
        module = infer_module_from_keyword_mapping(module_text, log_path_module_keywords, default=default_module)

        # 从 conf/config.py 的 TEST_CASE_DIRS 查表
        test_case_rel = ''
        env_key = self.get_env('env_key', '')
        project_root = project_dir
        if level1:
            if not project_dir.startswith('<') and '请配置' not in project_dir:
                project_root = project_dir
            else:
                project_root = os.path.dirname(os.path.dirname(os.path.dirname(SKILL_DIR)))
            sys.path.insert(0, project_root)
            from conf.config import Config
            for path_key in (level1, module_to_chinese(module)):
                test_case_rel = Config.TEST_CASE_DIRS.get(path_key, '')
                if test_case_rel:
                    break
            if not env_key:
                env_key = resolve_env(module, project_root)

        # 存完整绝对路径，step 文档直接用 {test_case_dir} 无需再拼接
        test_case_dir = os.path.join(project_root, test_case_rel)   # Full absolute path

        return {
            'env': self.env,
            'env_key': env_key,
            'presets': self.get_presets(module),
            'preset_fields': self.get_preset_fields(module),
            'menu': menu,
            'menu_url': url,
            'module': module,
            'func_code': func_code,
            'paths': {
                'project_dir': project_dir,
                'test_case_dir': test_case_dir,
            },
            'field_groups': [
                {
                    'name': '客户标识组',
                    'fields': ['客户标识类型', '客户标识'],
                    'order': ['客户标识类型', '客户标识'],
                    'aliases': {
                        '账户类型': '客户标识类型',
                        '账号': '客户标识',
                    },
                    'conditional': True,  # 仅在字段必填时无条件填入；非必填时可跳过（用于双路查询空查分支）
                }
            ],
        }

    def write_config_json(self, menu_name: str, output_dir: str = None) -> str:
        """将 Excel 三个 sheet 的数据写入 config.json，返回写入路径。

        Args:
            menu_name: 菜单名称
            output_dir: 输出目录。为 None 时写入 test_case_dir/{二级菜单}/{菜单名}/

        Returns:
            config.json 的完整路径
        """
        result = self.init_menu(menu_name)
        menu = result['menu']
        paths = result['paths']

        # 构建 config.json 结构
        env_data = dict(result['env'])
        env_data['env_key'] = result['env_key']

        config_data = {
            'env': env_data,
            'menu': menu,
            'presets': result['presets'],
            'preset_fields': result['preset_fields'],
            'field_groups': result.get('field_groups', []),
            'computed': {
                'menu_url': result['menu_url'],
                'module': result['module'],
                'func_code': result.get('func_code', ''),
                'test_case_dir': paths['test_case_dir'],
                'log_path_full': self._compute_log_path(result['module'], env_data),
            },
        }

        # 确定输出路径
        if output_dir is None:
            project_dir = self.get_env('project_dir', '')
            project_root = project_dir if (project_dir and not project_dir.startswith('<') and '请配置' not in project_dir) else os.path.dirname(os.path.dirname(os.path.dirname(SKILL_DIR)))
            test_case_dir = paths['test_case_dir']
            level2 = menu.get('导航二级菜单', '')
            menu_name_val = menu.get('菜单名称', '')
            output_dir = os.path.join(project_root, test_case_dir, level2, menu_name_val)

        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, 'config.json')
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(config_data, f, ensure_ascii=False, indent=2)

        return output_path

    def mark_done(self, menu_name: str) -> bool:
        """将指定菜单的状态标记为「已完成」，写回 Excel。

        Returns:
            True if marked, False if menu not found
        """
        from openpyxl import load_workbook as _lw
        wb = _lw(self.config_path)
        ws = wb['菜单导入']
        headers = [cell.value for cell in ws[1]]

        # 找到状态列和菜单名称列的索引
        status_idx = None
        name_idx = None
        for i, h in enumerate(headers):
            if h == '状态':
                status_idx = i
            elif h == '菜单名称':
                name_idx = i

        if status_idx is None or name_idx is None:
            print(f"[WARN] Excel 中缺少「状态」或「菜单名称」列", file=sys.stderr)
            return False

        found = False
        for row in ws.iter_rows(min_row=2):
            cell_name = row[name_idx].value
            if cell_name == menu_name:
                row[status_idx].value = '已完成'
                found = True
                break

        if found:
            wb.save(self.config_path)
            print(f"[OK] 已标记「{menu_name}」为已完成")
            # 清除缓存，下次读取生效
            self._import_menus = None
            self._menus = None
        else:
            print(f"[WARN] 未找到菜单: {menu_name}", file=sys.stderr)
        return found

    def to_dict(self) -> dict:
        return {
            'env': self.env,
            'menus': self.menus,
            'presets': self.presets,
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2)


# ========== 命令行查询 ==========
def main():
    config = ConfigReader()

    if len(sys.argv) < 2:
        print("用法: python config_reader.py <命令> [参数]")
        print("命令:")
        print("  env [key]          查看环境配置（指定key只看某项）")
        print("  init <菜单名>      一次性导出菜单全部配置（env+presets+menu+url）")
        print("  menus              列出未完成菜单")
        print("  menu <菜单名>      查看某菜单配置")
        print("  presets [key]      查看预埋数据（指定key只看某项）")
        print("  nav <菜单名>       生成导航代码")
        print("  menu_url <菜单名>  获取菜单URL路径")
        print("  case <菜单名> <类型> 生成hierarchical_case装饰器")
        print("  sql <菜单名>       生成预埋SQL路径代码")
        print("  base              列出全部基础配置表")
        print("  base_ta <code> <prd_type>  查询 tbtainfo 记录")
        print("  base_prd <table> <key>     查询产品表记录")
        print("  base_prds <table>          列出某产品表全部记录")
        print("  base_client <bank_acc>     查询客户映射")
        print("  compute_shard <module>     计算分片库（dxtrust/dxfund/fina），分表以日志物理表名为准")
        print("  sql_paths <菜单名>         输出 SQL 预埋文件的正确绝对路径")
        print("  defaults            查看探勘默认值（TA/产品/客户）")
        print("  mark_done <菜单名> 标记菜单为已完成")
        print("  match_presets <探勘JSON> 匹配探勘字段与预埋数据")
        print("  build_actions <form_fields.all JSON> 组装全选查询动作文件")
        print("  all                导出全部配置JSON")
        return

    cmd = sys.argv[1]

    if cmd == 'env':
        if len(sys.argv) > 2:
            val = config.get_env(sys.argv[2])
            print(f"{sys.argv[2]} = {val}")
        else:
            for k, v in config.env.items():
                print(f"  {k} = {v}")

    elif cmd == 'menus':
        for m in config.import_menus:
            print(f"  {m['菜单名称']}")

    elif cmd == 'init':
        if len(sys.argv) < 3:
            print("请指定菜单名"); return
        try:
            result = config.init_menu(sys.argv[2])
        except ValueError as e:
            print(f"[SKIP] {e}")
            return
        print(json.dumps(result, ensure_ascii=False, indent=2))
        # 同时写入 config.json
        config_path = config.write_config_json(sys.argv[2])
        print(f"\nconfig.json 已写入: {config_path}")

    elif cmd == 'menu':
        if len(sys.argv) < 3:
            print("请指定菜单名"); return
        menu = config.get_menu(sys.argv[2])
        if menu:
            for k, v in menu.items():
                print(f"  {k}: {v}")
        else:
            print(f"未找到菜单: {sys.argv[2]}")

    elif cmd == 'presets':
        if len(sys.argv) > 2:
            val = config.get_preset(sys.argv[2])
            print(f"{sys.argv[2]} = {val}")
        else:
            for k, v in config.presets.items():
                print(f"  {k} = {v}")

    elif cmd == 'nav':
        if len(sys.argv) < 3:
            print("请指定菜单名"); return
        print(config.gen_navigate_code(sys.argv[2]))

    elif cmd == 'menu_url':
        if len(sys.argv) < 3:
            print("请指定菜单名"); return
        url = config.get_menu_url(sys.argv[2])
        if url:
            print(url)
        else:
            print(f"未找到菜单URL: {sys.argv[2]}")

    elif cmd == 'case':
        if len(sys.argv) < 4:
            print("用法: python config_reader.py case <菜单名> <测试类型>"); return
        print(config.gen_hierarchical_case(sys.argv[2], sys.argv[3]))

    elif cmd == 'sql':
        if len(sys.argv) < 3:
            print("请指定菜单名"); return
        print(config.gen_preset_sql_code(sys.argv[2]))

    elif cmd == 'base_ta':
        if len(sys.argv) < 4:
            print("用法: python config_reader.py base_ta <ta_code> <prd_type(0|8)>"); return
        ta_code = sys.argv[2]
        prd_type = sys.argv[3]
        record = config.get_ta_record(ta_code, prd_type)
        if record:
            print(json.dumps(record, ensure_ascii=False, indent=2))
        else:
            print(f"未找到 tbtainfo 记录: ta_code={ta_code}, prd_type={prd_type}")

    elif cmd == 'base_prd':
        if len(sys.argv) < 4:
            print("用法: python config_reader.py base_prd <tbproduct|tbdxtrustproduct> <record_key>"); return
        table_name = sys.argv[2]
        record_key = sys.argv[3]
        record = config.get_product_record(table_name, record_key)
        if record:
            print(json.dumps(record, ensure_ascii=False, indent=2))
        else:
            print(f"未找到产品记录: {table_name}.{record_key}")

    elif cmd == 'base_prds':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py base_prds <tbproduct|tbdxtrustproduct>"); return
        table_name = sys.argv[2]
        records = config.get_products_by_module(table_name)
        if records:
            for k, v in records.items():
                if k.startswith('_'):
                    continue
                print(f"  {k}: prd_code={v.get('prd_code')}, prd_name={v.get('prd_name')}, prd_type={v.get('prd_type')}")
        else:
            print(f"未找到产品表: {table_name}")

    elif cmd == 'base_client':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py base_client <bank_acc>"); return
        record = config.get_client_mapping(sys.argv[2])
        if record:
            print(json.dumps(record, ensure_ascii=False, indent=2))
        else:
            print(f"未找到客户映射: {sys.argv[2]}")

    elif cmd == 'base':
        bt = config.base_tables
        for table_name, table_data in bt.items():
            if table_name.startswith('_'):
                continue
            records = table_data.get('records', {})
            count = len([k for k in records if not k.startswith('_')])
            print(f"  {table_name}: {count} 条记录")

    elif cmd == 'compute_shard':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py compute_shard <module>")
            print("  module: dxtrust | dxfund | fina")
            return
        module = sys.argv[2]
        try:
            result = config.compute_shard(module)
            if '--json' in sys.argv:
                print(json.dumps(result, ensure_ascii=False, indent=2))
            else:
                print(f"  customer_id = {result['customer_id']}")
                print(f"  in_client_no = {result['in_client_no']}")
                print(f"  last8 = {result['last8']}")
                print(f"  双库 = {'是' if result['has_db2'] else '否'}")
                print(f"  db_name = {result['db_name']}")
                print(f"  交易分表 = {result['trans_table_num']}")
                print(f"  公共分表 = {result['pub_table_num']}")
                print(f"  计算过程 = {result['formula']}")
        except Exception as e:
            if '--json' in sys.argv:
                print(json.dumps({'error': str(e)}, ensure_ascii=False))
            else:
                print(f"错误: {e}")

    elif cmd == 'sql_paths':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py sql_paths <菜单名>"); return
        menu_name = sys.argv[2]
        menu = config.get_menu(menu_name)
        if not menu:
            print(json.dumps({'error': f'未找到菜单: {menu_name}'}, ensure_ascii=False))
            return
        level1 = menu.get('导航一级菜单', '')
        level2 = menu.get('导航二级菜单', '')
        url = config.get_menu_url(menu_name)
        func_code = config._extract_func_code(url)

        project_dir = config.get_env('project_dir', '')
        project_root = project_dir if (project_dir and not project_dir.startswith('<') and '请配置' not in project_dir) else os.path.dirname(os.path.dirname(os.path.dirname(SKILL_DIR)))
        sys.path.insert(0, project_root)
        from conf.config import Config as _Cfg
        sql_base = _Cfg.SQL_DATA_DIRS.get(level1, '')

        if not sql_base or not func_code:
            print(json.dumps({'error': f'无法确定路径: level1={level1}, func_code={func_code}'}, ensure_ascii=False))
            return

        base_sql = os.path.join(project_root, sql_base, level2, menu_name, f'{func_code}_base.sql')
        preset_sql = os.path.join(project_root, sql_base, level2, menu_name, f'{func_code}.sql')

        result = {
            'menu': menu_name,
            'level1': level1,
            'level2': level2,
            'func_code': func_code,
            'sql_data_base': sql_base,
            'base_sql': base_sql,
            'preset_sql': preset_sql,
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))

    elif cmd == 'defaults':
        defaults = config.base_tables.get('_defaults', {})
        for k, v in defaults.items():
            if not k.startswith('_'):
                print(f"  {k} = {v}")

    elif cmd == 'mark_done':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py mark_done <菜单名>"); return
        config.mark_done(sys.argv[2])

    elif cmd == 'match_presets':
        if len(sys.argv) < 3:
            print("用法: python config_reader.py match_presets <探勘结果JSON文件>")
            print("  将探勘字段与 presets 匹配，输出 form_fields.all")
            return
        with open(sys.argv[2], 'r', encoding='utf-8') as f:
            explored = json.load(f)
        # 🔧 修复：从 config.json 读取 module，按模块过滤 preset_fields
        import glob
        config_jsons = glob.glob(os.path.join(os.path.dirname(sys.argv[2]), 'config.json'))
        module = ''
        if config_jsons:
            with open(config_jsons[0], 'r', encoding='utf-8') as f:
                cfg = json.load(f)
                module = cfg.get('computed', {}).get('module', '')
        if module:
            preset_fields = config.get_preset_fields(module)
            from stagehand_explorer.src.explore_wrapper import generate_field_values
            result = generate_field_values(explored, preset_fields)
        else:
            result = config.match_presets(explored)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if result['unmatched']:
            print(f"\n[WARN] 以下字段缺少预埋值，请在 Excel「预埋数据」sheet 中补充：", file=sys.stderr)
            for u in result['unmatched']:
                print(f"  - {u['label']} ({u['type']})", file=sys.stderr)

    elif cmd == 'build_actions':
        # 从 stdin 或文件读取 form_fields.all，输出 cdp_executor 动作文件
        if len(sys.argv) < 3:
            print("用法: python config_reader.py build_actions <form_fields.all JSON文件>")
            print("  输出全选查询动作文件（可直接传给 cdp_executor -a）")
            return
        with open(sys.argv[2], 'r', encoding='utf-8') as f:
            all_fields = json.load(f)
        actions = []
        for field in all_fields:
            label = field['label']
            ftype = field['type']
            value = field.get('value', '')
            if ftype == 'input':
                if value:
                    actions.append({'action': 'fill', 'label': label, 'value': value})
            elif ftype == 'single-select':
                actions.append({'action': 'select', 'label': label, 'value': value or '__FIRST_OPTION__'})
            elif ftype == 'multi-select':
                if value:
                    actions.append({'action': 'multi_select', 'label': label, 'values': [value]})

        # 依赖排序：级联字段（如 TA代码→产品代码）必须先选前置字段
        _DEPENDENCIES = {"产品代码": ["TA代码"], "账号": ["账户类型"]}
        for dep_label, prereqs in _DEPENDENCIES.items():
            norm = lambda value: ''.join(str(value or '').split()).lower()
            dep_idx = next((i for i, a in enumerate(actions) if norm(a.get('label')) == norm(dep_label)), None)
            if dep_idx is None:
                continue
            prereq_keys = {norm(p) for p in prereqs}
            prereq_indices = [i for i, a in enumerate(actions) if norm(a.get('label')) in prereq_keys]
            max_prereq = max(prereq_indices) if prereq_indices else -1
            if max_prereq > dep_idx:
                # 把依赖字段移到所有前置字段之后
                dep_action = actions.pop(dep_idx)
                actions.insert(max_prereq + 1, dep_action)

        actions.append({'action': 'click', 'label': '查询'})
        print(json.dumps(actions, ensure_ascii=False, indent=2))

    elif cmd == 'all':
        print(config.to_json())

    else:
        print(f"未知命令: {cmd}")


if __name__ == '__main__':
    main()
