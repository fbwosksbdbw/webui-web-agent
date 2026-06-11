"""生成测试配置Excel模板"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

current_year = datetime.now().year
wb = Workbook()

# ============ 通用样式 ============
header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
header_fill = PatternFill('solid', fgColor='4472C4')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_font = Font(name='Arial', size=10)
cell_align = Alignment(vertical='center', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
note_font = Font(name='Arial', size=9, color='808080', italic=True)
note_fill = PatternFill('solid', fgColor='FFF2CC')

def style_header(ws, row, max_col):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

def style_data(ws, row, max_col, is_note=False):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = note_font if is_note else cell_font
        cell.fill = note_fill if is_note else PatternFill()
        cell.alignment = cell_align
        cell.border = thin_border

# ============ Sheet1: 环境配置 ============
ws1 = wb.active
ws1.title = '环境配置'
headers1 = ['配置项', '配置键', '配置值', '说明', '示例值']
for i, h in enumerate(headers1, 1):
    ws1.cell(row=1, column=i, value=h)
style_header(ws1, 1, len(headers1))

env_data = [
    ['登录地址', 'login_url', 'http://10.20.150.149:10622/console/#/login', '系统登录页面URL', 'http://xxx:port/console/#/login'],
    ['应用服务器地址', 'app_server_host', '10.20.150.149', 'SSH连接的IP地址', '10.20.xxx.xxx'],
    ['应用服务器SSH端口', 'app_server_port', '10621', 'SSH端口（非默认22）', '22'],
    ['应用服务器用户', 'app_server_user', 'root', 'SSH登录用户名', 'root'],
    ['应用服务器密码', 'app_server_password', 'handsome', 'SSH登录密码', 'xxx'],
    ['日志路径', 'log_path', '/home/lcpt60/lcpt_jar/lcpt-web-manager-dxtrust/logs/', '应用日志目录', '/path/to/logs/'],
    ['日志文件名', 'log_file', 'lcpt-manager.log', '主日志文件名', 'app.log'],
    ['数据库主机', 'db_host', '10.20.29.245', '数据库服务器IP', '10.20.xxx.xxx'],
    ['数据库端口', 'db_port', '3306', '数据库端口', '3306'],
    ['数据库用户', 'db_user', 'bta60', '数据库连接用户名', 'db_user'],
    ['数据库密码', 'db_password', 'handsome', '数据库连接密码', 'xxx'],
    ['业务配置库', 'db_bc', 'ui_bc', '业务配置库名（菜单、TA信息等）', 'ui_bc'],
    ['业务库1', 'db_biz1', 'ui_fund1', '主业务库（含tbdxfund*表）', 'ui_fund1'],
    ['业务库2', 'db_biz2', 'ui_fund2', '历史表库（含tbyljdxfund*表）', 'ui_fund2'],
    ['登录用户', 'login_user', 'leily02', '默认登录用户名', 'user01'],
    ['登录密码', 'login_password', '11111111', '默认登录密码', '11111111'],
    ['环境标识', 'env_key', 'ifm60-dxfund', 'pytest --env=xxx 的值', 'ifm60-dxfund'],
    ['pytest命令路径', 'pytest_path', '<请配置您的 pytest 路径>', 'pytest执行命令（必填，如 .venv\\Scripts\\python.exe -m pytest）', 'python -m pytest'],
    ['Node.js 路径', 'node_path', '', 'Node.js 可执行文件路径（留空则自动搜索 PATH）', '/usr/local/bin/node'],
    ['项目根目录', 'project_dir', '<请配置您的项目根目录>', '项目根路径（必填）', '/path/to/project'],
    ['测试用例目录', 'testcases_dir', 'testcases/综合理财平台6.0/基金代销', '测试用例相对路径', 'testcases/xxx/xxx'],
    ['数据预埋目录', 'sql_data_dir', '数据预埋', '预埋SQL子目录名', '数据预埋'],
    ['测试报告目录', 'report_dir', '测试报告', '报告输出子目录名', '测试报告'],
]
for r, row_data in enumerate(env_data, 2):
    for c, val in enumerate(row_data, 1):
        ws1.cell(row=r, column=c, value=val)
    style_data(ws1, r, len(headers1))

ws1.column_dimensions['A'].width = 20
ws1.column_dimensions['B'].width = 22
ws1.column_dimensions['C'].width = 55
ws1.column_dimensions['D'].width = 35
ws1.column_dimensions['E'].width = 35

# ============ Sheet2: 菜单配置 ============
ws2 = wb.create_sheet('菜单配置')
headers2 = ['菜单名称', '功能代码', '导航一级菜单', '导航二级菜单', '导航三级菜单', '数据库表', '跨库标志', '预埋SQL文件名', '全量查询支持', '备注']
for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)
style_header(ws2, 1, len(headers2))

menu_data = [
    ['基金代销产品信息查询', 'ifmCCxglDxfundProductDetailQry', '基金代销', '查询管理', '基金代销产品信息查询', 'tbdxfundproductdetail', '否', 'ifmCCxglDxfundProductDetailQry.sql', '是', ''],
    ['基金代销历史账户类请求流水查询', 'ifmCXtglDxfundHisClientReqQry', '基金代销', '查询管理', '基金代销历史账户类请求流水查询', 'tbdxfundhisclientreq', '否', 'ifmCXtglDxfundHisClientReqQry.sql', '是', '双TA(JS+HT)'],
    ['基金代销历史冻结解冻流水查询', 'ifmCCxglDxfundHisFreezeTransReq', '基金代销', '查询管理', '基金代销历史冻结解冻流水查询', 'tbdxfundhisfrozen', '是(ui_fund1)', 'ifmCCxglDxfundHisFreezeTransReq.sql', '待验证', '单TA，交易代码6位'],
    ['基金代销交易类请求流水查询', 'ifmCCxglDxfundTransReqQuery', '基金代销', '查询管理', '基金代销交易类请求流水查询', 'tbdxfundtransreq,tbdxfundhistransreq', '否', 'ifmCCxglDxfundTransReqQuery.sql', '是', '多选TA，客户标识不可用'],
    ['基金代销历史交易类请求流水查询', 'ifmCXtglDxfundHiTransReqQuery', '基金代销', '查询管理', '基金代销历史交易类请求流水查询', 'tbdxfundhistransreq', '否', 'ifmCXtglDxfundHiTransReqQuery.sql', '是(需改日期)', '多选TA，默认日期需改'],
    ['基金代销客户份额查询', 'ifmCXtglDxfundShareInfoQuery', '基金代销', '查询管理', '基金代销客户份额查询', 'tbdxfundshareqry,tbdxfundshareextendqry', '否', 'ifmCXtglDxfundShareInfoQuery.sql', '是', '单TA，客户标识不可用'],
    ['基金代销交易类确认流水查询', 'ifmCXtglDxfundTransCfmQuery', '基金代销', '查询管理', '基金代销交易类确认流水查询', 'tbdxfundtranscfmqry', '否', 'ifmCXtglDxfundTransCfmQuery.sql', '否', '数据问题待修'],
]
for r, row_data in enumerate(menu_data, 2):
    for c, val in enumerate(row_data, 1):
        ws2.cell(row=r, column=c, value=val)
    style_data(ws2, r, len(headers2))

for col_letter, width in [('A', 30), ('B', 35), ('C', 14), ('D', 14), ('E', 30), ('F', 40), ('G', 15), ('H', 40), ('I', 15), ('J', 25)]:
    ws2.column_dimensions[col_letter].width = width

# ============ Sheet3: 预埋数据 ============
ws3 = wb.create_sheet('预埋数据')
headers3 = ['字段名', '默认值', '说明', '适用模块']
for i, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=i, value=h)
style_header(ws3, 1, len(headers3))

preset_data = [
    ['TA代码1', 'JS', '主TA代码', '所有菜单'],
    ['TA名称1', '嘉实基金', '主TA名称', '所有菜单'],
    ['TA下拉框选项1', 'JS:嘉实基金', '下拉框选择用', '所有菜单'],
    ['TA代码2', 'HT', '副TA代码（多选用）', '多选TA菜单'],
    ['TA名称2', 'HT基金', '副TA名称', '多选TA菜单'],
    ['TA下拉框选项2', 'HT:HT基金', '下拉框选择用', '多选TA菜单'],
    ['产品代码', '11', '产品信息查询用产品代码', '产品信息查询'],
    ['产品名称', 'AI自动化测试产品', '产品名称', '产品信息查询'],
    ['产品收益类型', '0:保证收益类', '产品收益类型下拉选项', '产品信息查询'],
    ['客户标识', 'A3060210020001', '统一客户标识', '所有菜单(部分不可用)'],
    ['账户类型', '0:入账账号', '账户类型下拉选项', '涉及账户类型的菜单'],
    ['客户标识类型', '0:入账账号', '客户标识类型下拉选项', '涉及客户标识类型的菜单'],
    ['查询日期起始', f'{current_year}-01-01', '历史类查询起始日期', '历史类查询菜单'],
    ['查询日期结束', f'{current_year}-12-31', '历史类查询结束日期', '历史类查询菜单'],
    ['冻结流水号', f'FZ{current_year}04280001', '冻结解冻用流水号', '冻结解冻查询'],
    ['解冻流水号', f'UF{current_year}04280001', '冻结解冻用解冻流水号', '冻结解冻查询'],
    ['交易代码(页面值)', '110222:基金代销份额冻结', '页面上交易代码下拉框选项', '冻结解冻查询'],
    ['交易代码(数据库值)', '022', '数据库中存储的交易代码', '冻结解冻查询'],
    ['冻结状态', '1:冻结成功', '冻结状态下拉选项', '冻结解冻查询'],
    ['冻结原因', '0:司法冻结', '冻结原因下拉选项', '冻结解冻查询'],
    ['是否有份额', '1:是', '是否有份额下拉选项', '客户份额查询'],
]
for r, row_data in enumerate(preset_data, 2):
    for c, val in enumerate(row_data, 1):
        ws3.cell(row=r, column=c, value=val)
    style_data(ws3, r, len(headers3))

for col_letter, width in [('A', 22), ('B', 35), ('C', 30), ('D', 25)]:
    ws3.column_dimensions[col_letter].width = width

# ============ Sheet4: 页面验证 ============
ws4 = wb.create_sheet('页面验证')
headers4 = ['菜单名称', 'TA下拉框类型', '表单字段数', '默认日期需修改', '客户标识过滤可用', '操作按钮选择器', '弹窗标题', '弹窗字段数', '特殊说明']
for i, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=i, value=h)
style_header(ws4, 1, len(headers4))

page_data = [
    ['基金代销产品信息查询', '单选', 6, '否', '待验证', '.h-omc-table-btn:visible', '控制标志信息', 10, '产品代码是下拉框不是输入框'],
    ['基金代销历史账户类请求流水查询', '多选', 6, '否', '待验证', '.h-omc-table-btn:visible', '预览信息', 15, '双TA查询'],
    ['基金代销历史冻结解冻流水查询', '单选', 13, '否', '是', '.h-omc-table-btn:visible', '详细信息', 27, '交易代码6位业务码，产品代码格式prd_code:ta_name'],
    ['基金代销交易类请求流水查询', '多选', 4, '否', '否(不可用)', '.h-omc-table-btn:visible', '预览信息', 20, '多选TA，客户标识过滤不可用，选完JS后需搜索选HT'],
    ['基金代销历史交易类请求流水查询', '多选', 18, f'是(2019→{current_year})', '待验证', '.h-omc-table-btn:visible', '预览信息', 25, '默认日期范围不含预埋数据，需改日期'],
    ['基金代销客户份额查询', '单选', 10, '否', '否(不可用)', '.h-omc-table-btn:visible', '详细信息', 35, '客户标识不可用，多条件用TA+是否有份额'],
    ['基金代销交易类确认流水查询', '多选', 23, f'是(2019→{current_year})', '待验证', '.h-omc-table-btn:visible', '详细信息', 30, '数据问题待修，后端查0条'],
]
for r, row_data in enumerate(page_data, 2):
    for c, val in enumerate(row_data, 1):
        ws4.cell(row=r, column=c, value=val)
    style_data(ws4, r, len(headers4))

for col_letter, width in [('A', 30), ('B', 12), ('C', 12), ('D', 18), ('E', 15), ('F', 28), ('G', 14), ('H', 12), ('I', 40)]:
    ws4.column_dimensions[col_letter].width = width

# ============ Sheet5: 字段映射 ============
ws5 = wb.create_sheet('字段映射')
headers5 = ['菜单名称', '字段名', '数据库存储值', '页面下拉框选项', '映射说明']
for i, h in enumerate(headers5, 1):
    ws5.cell(row=1, column=i, value=h)
style_header(ws5, 1, len(headers5))

mapping_data = [
    ['通用', '账户类型', '0', '0:入账账号', ''],
    ['通用', '客户标识类型', '0', '0:入账账号', ''],
    ['通用', '产品收益类型', '0', '0:保证收益类', ''],
    ['通用', '产品收益类型', '1', '1:保本浮动收益类', ''],
    ['通用', '产品收益类型', '2', '2:非保本浮动收益类', ''],
    ['冻结解冻查询', '交易代码', '022', '110222:基金代销份额冻结', '数据库存短码，页面用6位业务代码'],
    ['冻结解冻查询', '交易代码', '023', '110223:基金代销份额解冻', '同上'],
    ['冻结解冻查询', '冻结解冻状态', '1', '1:冻结成功', ''],
    ['冻结解冻查询', '冻结原因', '0', '0:司法冻结', ''],
    ['冻结解冻查询', '产品代码', '11', '11:嘉实基金', '选项格式 prd_code:ta_name'],
    ['客户份额查询', '是否有份额', '1', '1:是', ''],
    ['客户份额查询', '是否有份额', '0', '0:否', ''],
    ['', '', '', '', '← 新菜单的字段映射请追加到下方'],
]
for r, row_data in enumerate(mapping_data, 2):
    for c, val in enumerate(row_data, 1):
        ws5.cell(row=r, column=c, value=val)
    is_note = (r == len(mapping_data) + 1)
    style_data(ws5, r, len(headers5), is_note=is_note)

for col_letter, width in [('A', 25), ('B', 18), ('C', 20), ('D', 35), ('E', 35)]:
    ws5.column_dimensions[col_letter].width = width

# ============ 保存 ============
import os
SKILL_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
output_path = os.path.join(SKILL_DIR, 'config', 'test_config.xlsx')
wb.save(output_path)
print(f'配置模板已生成: {output_path}')
